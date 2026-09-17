"""Résumé text extraction.

Two outputs, both needed downstream:

* `text` - what an ATS would most likely see, used for keyword and semantic
  scoring.
* `structure` - layout facts (tables, images, column gutters, text boxes, header
  content) that the deterministic ATS checker in `ats_check.py` reasons over.

Extraction is deliberately kept separate from judgement: nothing here decides
whether a document is *good*, it only reports what is in it.
"""

from __future__ import annotations

import logging
import re
from dataclasses import dataclass, field
from pathlib import Path

logger = logging.getLogger(__name__)

SUPPORTED_SUFFIXES = {".pdf", ".docx", ".txt", ".md"}


@dataclass
class PageFacts:
    index: int
    width: float
    height: float
    char_count: int
    word_count: int
    ruled_table_count: int
    image_count: int
    # Fraction of page width taken by the widest vertical whitespace gutter that
    # runs through the body of the page, plus where its centre sits (0.0-1.0).
    gutter_width_ratio: float = 0.0
    gutter_center_ratio: float = 0.0
    header_text: str = ""
    footer_text: str = ""


@dataclass
class ExtractedDoc:
    text: str
    kind: str  # pdf | docx | txt
    filename: str
    ok: bool = True
    error: str = ""
    pages: list[PageFacts] = field(default_factory=list)
    # DOCX-only structural facts
    docx_table_count: int = 0
    docx_textbox_count: int = 0
    docx_inline_shape_count: int = 0
    docx_column_count: int = 1
    docx_header_footer_text: str = ""
    docx_uses_heading_styles: bool = False
    # Word list styles put the bullet glyph outside the text run, so a DOCX with
    # perfectly good bullets extracts as unmarked lines. Counting the styled
    # paragraphs recovers what the text alone cannot show.
    docx_list_paragraphs: int = 0

    @property
    def page_count(self) -> int:
        return max(len(self.pages), 1)

    @property
    def word_count(self) -> int:
        return len(self.text.split())

    @property
    def char_count(self) -> int:
        return len(self.text)


def extract(path: str | Path, *, original_filename: str = "") -> ExtractedDoc:
    path = Path(path)
    name = original_filename or path.name
    suffix = Path(name).suffix.lower() or path.suffix.lower()

    if suffix == ".pdf":
        return _extract_pdf(path, name)
    if suffix == ".docx":
        return _extract_docx(path, name)
    if suffix in {".txt", ".md"}:
        return ExtractedDoc(
            text=_normalise(path.read_text(encoding="utf-8", errors="replace")),
            kind="txt",
            filename=name,
        )
    return ExtractedDoc(
        text="",
        kind=suffix.lstrip(".") or "unknown",
        filename=name,
        ok=False,
        error=(
            f"Unsupported file type {suffix!r}. Compass reads .pdf, .docx, .txt "
            "and .md. Anything else is also a problem for most ATS parsers - "
            "export to PDF or DOCX first."
        ),
    )


# --------------------------------------------------------------------------- PDF


def _extract_pdf(path: Path, name: str) -> ExtractedDoc:
    try:
        import pdfplumber
    except ImportError:  # pragma: no cover
        return ExtractedDoc(
            text="", kind="pdf", filename=name, ok=False,
            error="pdfplumber is not installed. Run: pip install -r requirements.txt",
        )

    chunks: list[str] = []
    pages: list[PageFacts] = []
    try:
        with pdfplumber.open(str(path)) as pdf:
            for i, page in enumerate(pdf.pages):
                page_text = page.extract_text() or ""
                chunks.append(page_text)
                words = page.extract_words() or []

                try:
                    ruled_tables = len(page.find_tables())
                except Exception:  # pdfplumber can throw on malformed pages
                    ruled_tables = 0

                facts = PageFacts(
                    index=i,
                    width=float(page.width or 0),
                    height=float(page.height or 0),
                    char_count=len(page_text),
                    word_count=len(words),
                    ruled_table_count=ruled_tables,
                    image_count=len(page.images or []),
                )
                gutter_ratio, gutter_center = _detect_gutter(words, facts.width)
                facts.gutter_width_ratio = gutter_ratio
                facts.gutter_center_ratio = gutter_center
                facts.header_text = _band_text(words, facts.height, 0.0, 0.06)
                facts.footer_text = _band_text(words, facts.height, 0.94, 1.0)
                pages.append(facts)
    except Exception as exc:
        logger.exception("PDF extraction failed for %s", name)
        return ExtractedDoc(
            text="", kind="pdf", filename=name, ok=False,
            error=f"Could not read this PDF: {exc}",
        )

    return ExtractedDoc(
        text=_normalise("\n".join(chunks)), kind="pdf", filename=name, pages=pages
    )


def _detect_gutter(words: list[dict], page_width: float) -> tuple[float, float]:
    """Find the widest vertical whitespace channel with real content on both sides.

    A single-column résumé has no such channel. A two-column one has an obvious
    one, and that layout is the classic way an ATS scrambles a résumé into
    nonsense - so it is worth detecting even when there is no ruled table to
    give it away.
    """
    if not words or page_width <= 0:
        return 0.0, 0.0

    bins = 100
    occupied = [False] * bins
    step = page_width / bins
    for word in words:
        try:
            start = max(0, min(bins - 1, int(word["x0"] / step)))
            end = max(0, min(bins - 1, int(word["x1"] / step)))
        except (KeyError, TypeError, ValueError):
            continue
        for b in range(start, end + 1):
            occupied[b] = True

    best_width = 0
    best_center = 0.0
    run_start: int | None = None
    for b in range(bins + 1):
        filled = occupied[b] if b < bins else True
        if not filled and run_start is None:
            run_start = b
        elif filled and run_start is not None:
            run_end = b  # exclusive
            width = run_end - run_start
            left_has_content = any(occupied[:run_start])
            right_has_content = any(occupied[run_end:])
            if left_has_content and right_has_content and width > best_width:
                best_width = width
                best_center = (run_start + run_end) / 2 / bins
            run_start = None

    return best_width / bins, best_center


def _band_text(words: list[dict], page_height: float, top_frac: float, bottom_frac: float) -> str:
    if not words or page_height <= 0:
        return ""
    lo, hi = page_height * top_frac, page_height * bottom_frac
    picked = [w.get("text", "") for w in words if lo <= float(w.get("top", -1)) < hi]
    return " ".join(picked).strip()


# -------------------------------------------------------------------------- DOCX

_W_NS = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def _extract_docx(path: Path, name: str) -> ExtractedDoc:
    try:
        import docx
    except ImportError:  # pragma: no cover
        return ExtractedDoc(
            text="", kind="docx", filename=name, ok=False,
            error="python-docx is not installed. Run: pip install -r requirements.txt",
        )

    try:
        document = docx.Document(str(path))
    except Exception as exc:
        logger.exception("DOCX extraction failed for %s", name)
        return ExtractedDoc(
            text="", kind="docx", filename=name, ok=False,
            error=f"Could not read this DOCX: {exc}",
        )

    lines = [p.text for p in document.paragraphs]

    # Table cell text is part of what a parser sees, even though the table
    # itself is a parseability problem - collect it so scoring is not skewed.
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                lines.append(" | ".join(cells))

    try:
        raw_xml = document.element.xml
    except Exception:
        raw_xml = ""
    textbox_count = raw_xml.count("txbxContent")

    # Text inside a text box is invisible to `document.paragraphs`, so pull it
    # in separately - otherwise a text-box résumé scores as nearly empty.
    if textbox_count:
        try:
            for node in document.element.iter(f"{_W_NS}txbxContent"):
                for t in node.iter(f"{_W_NS}t"):
                    if t.text:
                        lines.append(t.text)
        except Exception:
            pass

    column_count = 1
    try:
        for section in document.sections:
            cols = section._sectPr.find(f"{_W_NS}cols")
            if cols is not None:
                num = cols.get(f"{_W_NS}num")
                if num and int(num) > column_count:
                    column_count = int(num)
    except Exception:
        pass

    header_footer: list[str] = []
    try:
        for section in document.sections:
            for part in (section.header, section.footer):
                for p in part.paragraphs:
                    if p.text.strip():
                        header_footer.append(p.text.strip())
    except Exception:
        pass

    uses_headings = False
    list_paragraphs = 0
    for paragraph in document.paragraphs:
        style_name = (paragraph.style.name or "").lower() if paragraph.style else ""
        if style_name.startswith("heading"):
            uses_headings = True
        if "list" in style_name and paragraph.text.strip():
            list_paragraphs += 1

    try:
        inline_shapes = len(document.inline_shapes)
    except Exception:
        inline_shapes = 0

    return ExtractedDoc(
        text=_normalise("\n".join(lines)),
        kind="docx",
        filename=name,
        docx_table_count=len(document.tables),
        docx_textbox_count=textbox_count,
        docx_inline_shape_count=inline_shapes,
        docx_column_count=column_count,
        docx_header_footer_text=" ".join(header_footer),
        docx_uses_heading_styles=uses_headings,
        docx_list_paragraphs=list_paragraphs,
    )


# ------------------------------------------------------------------- utilities


def _normalise(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t]+\n", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()
